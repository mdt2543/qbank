#!/usr/bin/env python3
"""
qbank_ingest.py — turn a generated quiz document into a loadable question bank.

Reads a .docx, .md, or .txt quiz, extracts questions/choices/keys/rationales,
runs quality checks, and writes CSV plus (optionally) Moodle GIFT and
Supabase SQL.

    python qbank_ingest.py quiz.docx --bank "Week 4 Histology"
    python qbank_ingest.py quiz.docx --bank "Week 4 Histology" --balance
    python qbank_ingest.py quiz.docx --bank "Week 4" --sql --gift

Outputs land next to the input unless --outdir is given.

Formats understood
------------------
Questions:
    Question 12
    (Objective 13478) The stem text...
        A. First choice
        B. Second choice
  Also: "12." or "Q12" as the number line; objective code optional and may
  be numeric or hyphenated (4860, 13478, 04-3237).

Answer keys, either in a later section or inline after each question:
    Question 12 — Correct Answer: Option B
    Mapped Objective Code: 13478
    Rationale: Because...
  Also: "Answer: B", "Correct answer: B", "Explanation:" for rationale.

Requires python-docx only for .docx input:  pip install python-docx
"""

import argparse
import collections
import csv
import pathlib
import random
import re
import sys

LETTERS = "ABCDE"


# --------------------------------------------------------------------------
# input
# --------------------------------------------------------------------------

def load_text(path):
    p = pathlib.Path(path)
    if not p.exists():
        sys.exit(f"File not found: {path}")
    if p.suffix.lower() == ".docx":
        try:
            import docx
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError:
            sys.exit("Reading .docx needs python-docx:  pip install python-docx")
        from docx.oxml.ns import qn
        doc = docx.Document(str(p))
        lines, images = [], {}
        rels = doc.part.rels
        # Walk the body in true document order. Iterating doc.paragraphs and
        # doc.tables separately puts every table at the end, which lets a
        # trailing text box bleed into the final question's rationale.
        for child in doc.element.body.iterchildren():
            tag = child.tag.split("}")[-1]
            if tag == "p":
                for blip in child.findall(".//" + qn("a:blip")):
                    rid = blip.get(qn("r:embed"))
                    if rid in rels:
                        part = rels[rid].target_part
                        images[len(lines)] = (
                            pathlib.PurePosixPath(part.partname).name, part.blob)
                lines.append(Paragraph(child, doc).text)
            elif tag == "tbl":
                for row in Table(child, doc).rows:
                    for cell in row.cells:
                        lines.extend(cell.text.splitlines())
        return "\n".join(lines), images
    return p.read_text(encoding="utf-8-sig"), {}


def clean(t):
    """Strip markdown emphasis and normalise whitespace and dashes."""
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t or "", flags=re.S)
    t = re.sub(r"\*(.+?)\*", r"\1", t, flags=re.S)
    t = t.replace("\\'", "'").replace('\\"', '"')
    t = re.sub(r"\\(?=[^\\])", "", t)
    t = t.replace("---", "\u2014").replace("--", "\u2013")
    return re.sub(r"\s+", " ", t).strip()


# --------------------------------------------------------------------------
# patterns
# --------------------------------------------------------------------------

RE_QNUM = re.compile(r"^\s*(?:\*\*)?(?:Question|Q)?\s*(\d{1,3})\s*[.):\\]?\s*(?:\*\*)?\s*$", re.I)
RE_QINLINE = re.compile(r"^\s*(?:\*\*)?(?:Question|Q)\s*(\d{1,3})\s*[.):\\—–-]*\s*(.+)$", re.I)
# Bare "1." / "1)" numbering. Ambiguous with ordinary numbered lists, so it
# only counts when the number continues the sequence — see question_start().
RE_QBARE = re.compile(r"^\s*(?:\*\*)?(\d{1,3})\s*[.)]\s+(.+)$")


def question_start(line, last_num):
    """Return (number, inline_text) if this line begins a question, else None.

    `last_num` is the highest question number seen so far; bare numbering is
    accepted only when it continues the sequence, so a "1)" inside a stem or
    a numbered list doesn't get mistaken for a new question.
    """
    m = RE_QNUM.match(line)
    if m:
        return int(m.group(1)), ""
    m = RE_QINLINE.match(line)
    if m and not RE_CHOICE.match(line):
        return int(m.group(1)), m.group(2)
    m = RE_QBARE.match(line)
    if m and not RE_CHOICE.match(line):
        n = int(m.group(1))
        if n == (last_num or 0) + 1 and len(m.group(2)) > 20:
            return n, m.group(2)
    return None
RE_OBJ = re.compile(r"^\(\s*Objective\s*([\d][\d\-]*)\s*\)\s*", re.I)
RE_CHOICE = re.compile(r"^\s*>?\s*(?:\*\*)?([A-E])[.):]\s*(?:\*\*)?\s*(.+)$")
RE_KEY = re.compile(
    r"(?:Question\s*(\d{1,3}).*?)?"
    r"(?:Correct\s*Answer|Correct|Answer|Key)\s*[:\u2014\u2013-]*\s*(?:Option\s*)?([A-E])\b",
    re.I,
)
RE_MAPPED = re.compile(r"Mapped\s*Objective\s*Code\s*[:\-]?\s*([\d][\d\-]*)", re.I)
# A section header naming an objective, e.g. "05-1657 — Name the factors that…".
# Questions beneath it inherit that code until the next header.
RE_OBJHEADER = re.compile(r"^\s*([\d][\d\-]{2,})\s*[\u2014\u2013]{1,2}\s*(.{10,})$")
# A key line that starts with a bare question number, e.g. "12.  Answer: B — …".
RE_KEYNUM = re.compile(r"^\s*(?:\*\*)?(\d{1,3})\s*[.)]\s")
RE_RATIONALE = re.compile(r"^\s*>?\s*(?:\*\*)?(?:Rationale|Explanation|Why)\s*[:\-]\s*(?:\*\*)?\s*(.*)$", re.I)
RE_KEYSECTION = re.compile(r"(answer\s*key|rationales?|explanations?|part\s*(ii|2)\b)", re.I)
# Headings and boilerplate that should end a rationale rather than join it.
RE_STOP = re.compile(
    r"^\s*(instructions?|exam\s*overview|part\s+[ivx0-9]|appendix|end\s+of|"
    r"references?|notes?\s*:|disclaimer|source)\b", re.I)


# --------------------------------------------------------------------------
# parsing
# --------------------------------------------------------------------------

def split_sections(lines):
    """Find where an answer-key section begins, if there is one."""
    for i, ln in enumerate(lines):
        s = clean(ln)
        if not (len(s) < 80 and RE_KEYSECTION.search(s)):
            continue
        if RE_CHOICE.match(s) or RE_RATIONALE.match(s) or RE_KEY.search(s):
            continue  # a rationale or answer line, not a heading
        # Only a real divider if keys follow it and none appear before it.
        # Documents with inline answers have no separate key section.
        if RE_KEY.search("\n".join(lines[i:i + 60])) and not any(
                RE_KEY.search(clean(x)) for x in lines[:i]):
            return lines[:i], lines[i:]
    return lines, []


def parse_questions(lines, images_by_line=None):
    """Extract stems and choices. Returns {num: {...}}.

    images_by_line maps a line index to (filename, caption). A figure applies
    to every question from that point until the next objective header, which
    matches how these documents are written — one figure above a group, with
    the questions referring to "the figure above".
    """
    images_by_line = images_by_line or {}
    out = {}
    num = None
    pending_stem = []
    state = None  # 'stem' | 'choices'
    current_obj = None
    current_img = None

    def flush():
        if num is None:
            return
        raw = clean(" ".join(pending_stem))
        obj = None
        m = RE_OBJ.match(raw)
        if m:
            obj = m.group(1)
            raw = raw[m.end():]
        rec = out.setdefault(num, {"choices": {}, "obj": None, "stem": "",
                                   "image": None, "caption": None})
        if raw:
            rec["stem"] = raw
        if obj:
            rec["obj"] = obj

    for idx, raw_line in enumerate(lines):
        line = raw_line.rstrip()
        if idx in images_by_line:
            current_img = images_by_line[idx]
        if not line.strip():
            continue

        m = RE_OBJHEADER.match(clean(line))
        if m and not RE_CHOICE.match(clean(line)) and not RE_KEY.search(clean(line)):
            flush()
            current_obj = m.group(1)
            current_img = None
            num = None
            state = None
            pending_stem = []
            continue

        start = question_start(clean(line), max(out) if out else 0)
        if start:
            flush()
            num, inline = start
            pending_stem = [inline] if inline else []
            out.setdefault(num, {"choices": {}, "obj": None, "stem": "",
                                 "image": None, "caption": None})
            if current_obj and not out[num]["obj"]:
                out[num]["obj"] = current_obj
            if current_img and not out[num]["image"]:
                out[num]["image"], out[num]["caption"] = current_img
            state = "stem"
            continue

        m = RE_CHOICE.match(clean(line))
        if m and num is not None:
            if state == "stem":
                flush()
                pending_stem = []
            state = "choices"
            out[num]["choices"][m.group(1)] = clean(m.group(2))
            continue

        c = clean(line)
        # Inline answer/rationale lines end the choice block; they are picked
        # up by parse_keys, and must not be glued onto the last choice.
        if RE_KEY.search(c) or RE_RATIONALE.match(c) or RE_MAPPED.search(c):
            state = "afterchoices"
            continue

        if num is not None and state == "stem":
            pending_stem.append(line)
        elif num is not None and state == "choices":
            if out[num]["choices"]:
                last = sorted(out[num]["choices"])[-1]
                out[num]["choices"][last] += " " + c

    flush()
    return {n: r for n, r in out.items() if r["choices"]}


def parse_keys(lines, questions):
    """Extract answer letters, mapped objectives and rationales."""
    keys, objs, rats = {}, {}, {}
    keytext = {}
    current = None
    collecting = False
    buf = []

    def stash():
        if current is not None and buf:
            rats.setdefault(current, clean(" ".join(buf)))

    seen_max = 0

    for raw_line in lines:
        line = clean(raw_line)
        if not line:
            continue

        # Track question headings so an inline "Answer: B" attaches to the
        # question it follows, even with no number on the answer line.
        start = question_start(line, seen_max)
        if start and not RE_KEY.search(line):
            stash()
            buf = []
            collecting = False
            current = start[0]
            seen_max = max(seen_max, current)
            continue

        # Objective headers repeat inside the key section; they end any
        # rationale in progress but carry no answer of their own.
        if (RE_OBJHEADER.match(line) and not RE_CHOICE.match(line)
                and not RE_KEY.search(line)):
            stash()
            buf = []
            collecting = False
            continue

        m = RE_KEY.search(line)
        if m:
            stash()
            buf = []
            if m.group(1):
                current = int(m.group(1))
            else:
                bare = RE_KEYNUM.match(line)
                if bare:
                    current = int(bare.group(1))
                elif current is None:
                    current = seen_max + 1
                elif current in keys:
                    # A bare sequential list of answers with no numbering.
                    current += 1
            seen_max = max(seen_max, current or 0)
            if current is not None:
                keys[current] = m.group(2).upper()
                # Some documents restate the answer after a dash:
                # "12. Answer: B — the restated option text". Keep it so the
                # QA pass can confirm the letter points at the right option.
                tail = line[m.end():].strip()
                tm = re.match(r"^[\u2014\u2013:-]\s*(.{5,})$", tail)
                if tm:
                    keytext[current] = tm.group(1).strip()
            # Prose following an answer line is the rationale even when it
            # carries no "Rationale:" label. An explicit label overrides it.
            collecting = True
            continue

        m = RE_MAPPED.search(line)
        if m and current is not None:
            objs[current] = m.group(1)
            continue

        m = RE_RATIONALE.match(line)
        if m and current is not None:
            buf = []
            rats.pop(current, None)  # labelled text wins over unlabelled prose
            if m.group(1):
                buf = [m.group(1)]
            collecting = True
            continue

        if collecting and current is not None:
            if (RE_QNUM.match(line)
                    or re.match(r"^\s*Question\s*\d+", line, re.I)
                    or RE_STOP.match(line)):
                stash()
                buf = []
                collecting = False
            else:
                buf.append(line)

    stash()
    return keys, objs, rats, keytext


RE_FIGCUE = re.compile(
    r"\b(above|below|figure|shown|tracing|panel|image|graph|curve|loop|"
    r"spirogram|micrograph|slide|arrow)\b", re.I)


def assign_images(lines, images, questions):
    """Map each image to the question(s) that use it.

    Images sit above a block of questions as a shared stimulus. The one that
    actually refers to it ("the curve above…") gets it; if nothing in the
    block refers to a figure, the first question after the image does.
    """
    if not images:
        return {}, {}

    # Line number at which each question's text begins.
    qline, last = {}, 0
    for i, raw in enumerate(lines):
        st = question_start(clean(raw), last)
        if st and st[0] in questions:
            qline[st[0]] = i
            last = max(last, st[0])

    starts = sorted(qline.items(), key=lambda kv: kv[1])
    assigned, captions = {}, {}

    for img_line in sorted(images):
        after = [n for n, ln in starts if ln > img_line]
        if not after:
            continue
        # Questions between this image and the next one.
        nxt = min((l for l in images if l > img_line), default=10 ** 9)
        block = [n for n in after if qline[n] < nxt]
        if not block:
            continue
        cited = [n for n in block if RE_FIGCUE.search(questions[n]["stem"])]
        targets = cited or block[:1]

        # A caption is a short non-question line just after the image.
        caption = ""
        for j in range(img_line, min(img_line + 3, len(lines))):
            t = clean(lines[j])
            if t and not question_start(t, 0) and not RE_CHOICE.match(t):
                caption = t
                break

        for n in targets:
            assigned[n] = images[img_line]
            if caption:
                captions[n] = caption
    return assigned, captions


# --------------------------------------------------------------------------
# quality checks
# --------------------------------------------------------------------------

def qa_report(items, expected_objectives=None):
    problems, warnings, notes = [], [], []
    n = len(items)

    for it in items:
        ref = it["ref"]
        if not it["correct"]:
            problems.append(f"{ref}: no answer key found")
        elif it["correct"] not in it["choices"]:
            problems.append(
                f"{ref}: key '{it['correct']}' not among choices "
                f"({', '.join(sorted(it['choices']))})"
            )
        if len(it["choices"]) < 2:
            problems.append(f"{ref}: only {len(it['choices'])} choices parsed")
        if len(it["stem"]) < 15:
            problems.append(f"{ref}: stem looks truncated — {it['stem'][:50]!r}")
        if not it["explanation"]:
            warnings.append(f"{ref}: no rationale")
        if not it["obj"]:
            warnings.append(f"{ref}: no objective code")

        # If the document restates the answer text next to the letter, make
        # sure the letter actually points at that option. This catches a
        # mis-lettered key, which nothing else here would notice.
        kt = it.get("keytext")
        if kt and it["correct"] in it["choices"]:
            def _n(x):
                return re.sub(r"\s+", " ", x).strip().lower().rstrip(".")
            if _n(it["choices"][it["correct"]])[:60] != _n(kt)[:60]:
                problems.append(
                    f"{ref}: key says {it['correct']} but that option's text "
                    f"does not match the restated answer\n"
                    f"      option {it['correct']}: {it['choices'][it['correct']][:70]}\n"
                    f"      restated  : {kt[:70]}")

        bodies = [v.lower() for v in it["choices"].values()]
        dupes = [b for b, c in collections.Counter(bodies).items() if c > 1]
        if dupes:
            problems.append(f"{ref}: duplicate answer choices")

    counts = collections.Counter(len(it["choices"]) for it in items)
    if len(counts) > 1:
        notes.append("Mixed choice counts: " + ", ".join(
            f"{k} choices × {v}" for k, v in sorted(counts.items())))

    keyed = [it for it in items if it["correct"] in it["choices"]]
    dist = collections.Counter(it["correct"] for it in keyed)
    if keyed:
        top, topn = dist.most_common(1)[0]
        share = topn / len(keyed)
        even = 1 / max(len(set(l for it in keyed for l in it["choices"])), 1)
        if share > max(0.45, even * 1.8):
            warnings.append(
                f"Answer position skew: {top} is correct {topn}/{len(keyed)} "
                f"({share:.0%}). Use --balance, or shuffle at display time."
            )

    stems = collections.Counter(it["stem"][:90].lower() for it in items)
    for s, c in stems.items():
        if c > 1:
            warnings.append(f"{c} questions share a near-identical stem: {s[:60]}…")

    found = {it["obj"] for it in items if it["obj"]}
    if expected_objectives:
        missing = set(expected_objectives) - found
        extra = found - set(expected_objectives)
        if missing:
            problems.append("Objectives with no questions: " + ", ".join(sorted(missing)))
        if extra:
            warnings.append("Objectives not on your list: " + ", ".join(sorted(extra)))

    return problems, warnings, notes, dist, found


# --------------------------------------------------------------------------
# balancing
# --------------------------------------------------------------------------

def balance(items, seed):
    """Move each key so positions are used as evenly as possible."""
    rng = random.Random(seed)
    groups = collections.defaultdict(list)
    for it in items:
        groups[len(it["choices"])].append(it)

    for size, group in groups.items():
        slots = list(LETTERS[:size])
        targets = (slots * ((len(group) // size) + 1))[:len(group)]
        rng.shuffle(targets)
        for it, target in zip(group, targets):
            key = it["correct"]
            if key not in it["choices"]:
                continue
            correct_body = it["choices"][key]
            others = [v for k, v in sorted(it["choices"].items()) if k != key]
            rng.shuffle(others)
            new = {}
            it_others = iter(others)
            for L in slots:
                new[L] = correct_body if L == target else next(it_others)
            it["choices"] = new
            it["correct"] = target
    return items


# --------------------------------------------------------------------------
# output
# --------------------------------------------------------------------------

CSV_COLS = (["ref", "qbank", "topic", "difficulty", "stem"]
            + [f"choice_{c}" for c in "abcde"]
            + ["correct", "explanation", "image_path", "image_caption"]
            + [f"rationale_{c}" for c in "abcde"])


def write_csv(items, path, bank):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=CSV_COLS)
        w.writeheader()
        for it in items:
            row = {
                "ref": it["ref"], "qbank": bank,
                "topic": f"Objective {it['obj']}" if it["obj"] else "General",
                "difficulty": "", "stem": it["stem"],
                "correct": it["correct"], "explanation": it["explanation"],
                "image_path": it.get("image") or "",
                "image_caption": it.get("caption") or "",
            }
            for L in "abcde":
                row[f"choice_{L}"] = it["choices"].get(L.upper(), "")
                row[f"rationale_{L}"] = ""
            w.writerow(row)


GIFT_SPECIAL = re.compile(r"([~=#{}:\\])")


def esc_gift(t):
    return GIFT_SPECIAL.sub(r"\\\1", (t or "").strip())


def slugify(t):
    return re.sub(r"[^a-z0-9]+", "-", (t or "").lower()).strip("-") or "qbank"


def write_gift(items, path, bank):
    by_topic = collections.defaultdict(list)
    for it in items:
        by_topic[f"Objective {it['obj']}" if it["obj"] else "General"].append(it)

    lines = ["// Generated by qbank_ingest.py", "// Import via Question bank > Import > GIFT.", ""]
    for topic, rows in by_topic.items():
        lines += [f"$CATEGORY: $course$/top/{bank}/{topic}", ""]
        for it in rows:
            lines.append(f"::{esc_gift(it['ref'])}:: [html]{esc_gift(it['stem'])} {{")
            for L, body in sorted(it["choices"].items()):
                lines.append(f"    {'=' if L == it['correct'] else '~'}{esc_gift(body)}")
            if it["explanation"]:
                lines.append(f"    ####{esc_gift(it['explanation'])}")
            lines += ["}", ""]
    pathlib.Path(path).write_text("\n".join(lines), encoding="utf-8")


def esc_sql(t):
    if not t:
        return "NULL"
    return "'" + str(t).replace("'", "''") + "'"


def write_sql(items, path, bank):
    topics = sorted({f"Objective {it['obj']}" if it["obj"] else "General" for it in items})
    out = ["-- Generated by qbank_ingest.py. Safe to re-run.", "begin;", ""]
    out.append(f"insert into qbanks (slug, title) values "
               f"({esc_sql(slugify(bank))}, {esc_sql(bank)})\n"
               "  on conflict (slug) do update set title = excluded.title;")
    for t in topics:
        out.append(f"insert into topics (name) values ({esc_sql(t)})\n"
                   "  on conflict (name) do nothing;")
    out.append("")
    for it in items:
        topic = f"Objective {it['obj']}" if it["obj"] else "General"
        out.append(
            "insert into questions (external_ref, qbank_id, topic_id, stem, explanation,\n"
            "                       image_path, image_caption)\n"
            f"values ({esc_sql(it['ref'])},\n"
            f"        (select id from qbanks where slug = {esc_sql(slugify(bank))}),\n"
            f"        (select id from topics where name = {esc_sql(topic)}),\n"
            f"        {esc_sql(it['stem'])}, {esc_sql(it['explanation'])},\n"
            f"        {esc_sql(it.get('image'))}, {esc_sql(it.get('caption'))})\n"
            "on conflict (external_ref) do update set\n"
            "  qbank_id = excluded.qbank_id, topic_id = excluded.topic_id,\n"
            "  stem = excluded.stem, explanation = excluded.explanation,\n"
            "  image_path = excluded.image_path,\n"
            "  image_caption = excluded.image_caption;"
        )
        out.append("delete from choices where question_id = "
                   f"(select id from questions where external_ref = {esc_sql(it['ref'])});")
        for L, body in sorted(it["choices"].items()):
            out.append(
                "insert into choices (question_id, label, body, is_correct)\n"
                f"values ((select id from questions where external_ref = {esc_sql(it['ref'])}),\n"
                f"        {esc_sql(L)}, {esc_sql(body)}, "
                f"{'true' if L == it['correct'] else 'false'});"
            )
        out.append("")
    out += ["commit;", ""]
    out.append(f"-- publish when ready:")
    out.append(f"-- update qbanks set is_published = true where slug = {esc_sql(slugify(bank))};")
    pathlib.Path(path).write_text("\n".join(out), encoding="utf-8")


# --------------------------------------------------------------------------
# main
# --------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", help="quiz document (.docx, .md, .txt)")
    ap.add_argument("--bank", required=True, help='bank title, e.g. "Week 4 Histology"')
    ap.add_argument("--prefix", help="ref prefix (default: derived from bank)")
    ap.add_argument("--outdir", default=None)
    ap.add_argument("--balance", action="store_true",
                    help="redistribute correct answers evenly across positions")
    ap.add_argument("--seed", type=int, default=20260913)
    ap.add_argument("--gift", action="store_true", help="also write Moodle GIFT")
    ap.add_argument("--sql", action="store_true", help="also write Supabase SQL")
    ap.add_argument("--objectives", help="comma-separated codes you expect covered")
    ap.add_argument("--strict", action="store_true",
                    help="exit non-zero if any problems are found")
    args = ap.parse_args()

    src = pathlib.Path(args.input)
    outdir = pathlib.Path(args.outdir) if args.outdir else src.parent
    outdir.mkdir(parents=True, exist_ok=True)
    stem_name = slugify(args.bank)
    prefix = args.prefix or "".join(w[0] for w in re.findall(r"[A-Za-z]+", args.bank))[:4].upper() or "Q"

    text, raw_images = load_text(src)
    lines = text.splitlines()

    # Write figures out under stable names and work out each one's caption:
    # the next non-empty line that isn't a question, choice or objective
    # header. Names are positional, so re-running produces the same files.
    imgdir = outdir / f"{stem_name}-images"
    images_by_line = {}
    for n, (line_idx, (orig, blob)) in enumerate(sorted(raw_images.items()), 1):
        ext = pathlib.Path(orig).suffix or ".png"
        name = f"{stem_name}-fig-{n:02d}{ext}"
        imgdir.mkdir(parents=True, exist_ok=True)
        (imgdir / name).write_bytes(blob)
        caption = ""
        for look in lines[line_idx:line_idx + 4]:
            c = clean(look)
            if not c:
                continue
            if (question_start(c, 0) or RE_CHOICE.match(c) or RE_OBJHEADER.match(c)
                    or re.match(r"^Sources?\s*:", c, re.I)):
                break
            caption = c
            break
        images_by_line[line_idx] = (name, caption)

    if images_by_line:
        print(f"Extracted {len(images_by_line)} figure(s) to {imgdir}", file=sys.stderr)

    body, keysec = split_sections(lines)
    body_images = {i: v for i, v in images_by_line.items() if i < len(body)}

    questions = parse_questions(body, body_images)
    keys, objs, rats, keytext = parse_keys(keysec or body, questions)

    items = []
    for num in sorted(questions):
        q = questions[num]
        items.append({
            "num": num,
            "ref": f"{prefix}-{num:03d}",
            "stem": q["stem"],
            "obj": q["obj"] or objs.get(num),
            "choices": q["choices"],
            "correct": (keys.get(num) or "").upper(),
            "explanation": rats.get(num, ""),
            "keytext": keytext.get(num, ""),
            "image": q.get("image"),
            "caption": q.get("caption"),
        })

    if not items:
        sys.exit("No questions parsed. Check the document format against the header of this script.")

    print(f"\nParsed {len(items)} questions from {src.name}", file=sys.stderr)

    expected = [c.strip() for c in args.objectives.split(",")] if args.objectives else None
    problems, warnings, notes, dist, found = qa_report(items, expected)

    if args.balance:
        items = balance(items, args.seed)
        problems2, warnings2, notes2, dist, _ = qa_report(items, expected)
        warnings = [w for w in warnings2 if "skew" not in w]
        print("Balanced answer positions.", file=sys.stderr)

    print("\n── QA report " + "─" * 46, file=sys.stderr)
    print(f"Answer positions : " +
          "  ".join(f"{L}={dist.get(L, 0)}" for L in LETTERS if dist.get(L) or L in "ABCD"),
          file=sys.stderr)
    print(f"Objectives       : {len(found)} distinct", file=sys.stderr)
    for n in notes:
        print(f"note    · {n}", file=sys.stderr)
    for w in warnings:
        print(f"WARNING · {w}", file=sys.stderr)
    for p in problems:
        print(f"PROBLEM · {p}", file=sys.stderr)
    if not problems and not warnings:
        print("No issues found.", file=sys.stderr)
    print("─" * 60 + "\n", file=sys.stderr)

    csv_path = outdir / f"{stem_name}.csv"
    write_csv(items, csv_path, args.bank)
    print(f"wrote {csv_path}", file=sys.stderr)

    if args.gift:
        p = outdir / f"{stem_name}.gift.txt"
        write_gift(items, p, args.bank)
        print(f"wrote {p}", file=sys.stderr)

    if args.sql:
        p = outdir / f"{stem_name}.seed.sql"
        write_sql(items, p, args.bank)
        print(f"wrote {p}", file=sys.stderr)

    if problems and args.strict:
        sys.exit(1)


if __name__ == "__main__":
    main()